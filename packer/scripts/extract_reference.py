# -*- coding: utf-8 -*-
"""Extract the packer's ground truth from the BRC workbooks.

For every populated location sheet in both workbooks this reads the slot grid and
recovers:

  * the GROUP structure - a group is one evaluation domain, detected by its ZP
    (counting point) numbering restarting at 1. Backplanes belong wholly to one
    group; racks may span several.
  * the ACTUAL layout the planner drew - rack count, backplane mix, PSC, COM.

Writes packer/fixtures/reference.json, which score.ts reads.
"""
import openpyxl, json, os, re, io, warnings
warnings.filterwarnings('ignore')

HERE = os.path.dirname(os.path.abspath(__file__))
PACKER = os.path.dirname(HERE)
BASE = os.path.dirname(PACKER)
SRC = {'ABS':  os.path.join(BASE, 'BOM CAL', 'ABS V.1_2025-BRC with BD BOM.xlsm'),
       'Yard': os.path.join(BASE, 'BOM CAL', 'Yard V.1_2025-BRC with BD BOM.xlsm')}

BLOCK0, PERIOD, BLOCKS = 56, 14, 8      # first grid block, row period, blocks per sheet
COL0, COL1 = 3, 30
R_LABEL, R_BOARD, R_TE, R_ZP = 1, 2, 3, 4   # offsets from a block's base row

# Gesamt rows carrying the planner's own totals, for reconciliation.
GESAMT = {'racks': 14, 'psc': 35, 'com': 36, 'aeb': 39, 'io': 40,
          'BP-PWR-0': 16, 'BP-PWR-4': 20, 'BP-PWR-8': 22,
          'BP-EXB-1': 29, 'BP-EXB-2': 30, 'BP-EXB-4': 32}


def cell(ws, r, c):
    v = ws.cell(r, c).value
    return None if v is None or (isinstance(v, str) and not v.strip()) else v


def parse_sheet(ws):
    """Return (racks, groups). Each rack is a list of backplane dicts."""
    racks, backplanes = [], []
    for b in range(BLOCKS):
        base = BLOCK0 + b * PERIOD
        labels = [(c, cell(ws, base + R_LABEL, c)) for c in range(COL0, COL1)]
        if not any(v for _, v in labels):
            continue
        rack_code = next((str(v) for _, v in labels if v and str(v).startswith('BGT')), None)
        starts = [(c, str(v)) for c, v in labels if v and str(v).startswith('BP-')]
        if rack_code:
            racks.append({'index': len(racks) + 1, 'type': rack_code, 'backplanes': []})
        if not racks:
            continue
        rack = racks[-1]
        for i, (c0, code) in enumerate(starts):
            c1 = starts[i + 1][0] if i + 1 < len(starts) else COL1
            slots, zps, te = [], [], 0
            for c in range(c0, c1):
                board = cell(ws, base + R_BOARD, c)
                if board is None:
                    continue
                w = cell(ws, base + R_TE, c)
                te += w if isinstance(w, (int, float)) else 0
                z = cell(ws, base + R_ZP, c)
                if isinstance(z, (int, float)):
                    zps.append(int(z))
                slots.append(str(board).strip())
            bp = {'code': code, 'te': te, 'slots': slots, 'zp': zps,
                  'rack': rack['index'],
                  'aeb': slots.count('AEB'),
                  'io': sum(slots.count(t) for t in ('IO-EXB', 'CO-EXB')),
                  'com': sum(slots.count(t) for t in ('COM-AdC', 'COM-xxx')),
                  'psc': slots.count('PSC') + slots.count('PSC-R'),
                  'sparePsc': slots.count('spare-PSC')}
            rack['backplanes'].append(bp)
            backplanes.append(bp)

    # Group boundary: a backplane whose ZP numbering restarts at 1 after the
    # current group has already numbered a counting point.
    groups, cur, seen = [], [], False
    for bp in backplanes:
        if bp['zp'] and min(bp['zp']) == 1 and seen:
            groups.append(cur); cur, seen = [], False
        cur.append(bp)
        if bp['zp']:
            seen = True
    if cur:
        groups.append(cur)
    return racks, groups


def main():
    out = {}
    for proj, path in SRC.items():
        wv = openpyxl.load_workbook(path, data_only=True)
        g = wv['Gesamt']
        locs = []
        for idx in range(1, 31):
            name = cell(g, 4, 2 + idx)
            if not name or re.fullmatch(r'(Tabelle|Loc)\s*\d+', str(name).strip()):
                continue
            sheet = '%02d' % idx
            if sheet not in wv.sheetnames:
                continue
            racks, groups = parse_sheet(wv[sheet])
            if not racks:
                continue
            actual_bp = {}
            for r in racks:
                for bp in r['backplanes']:
                    actual_bp[bp['code']] = actual_bp.get(bp['code'], 0) + 1
            gz = {k: (cell(g, row, 2 + idx) or 0) for k, row in GESAMT.items()}
            locs.append({
                'id': 'L%02d' % len(locs) and 'L%02d' % (len(locs) + 1),
                'name': str(name).strip(), 'sheet': sheet,
                'groups': [{
                    'aeb': sum(b['aeb'] for b in grp),
                    'ioExb': sum(b['io'] for b in grp),
                    'com': sum(b['com'] for b in grp),
                    # PSC actually equipped in this group's power slots. Not
                    # derivable - the planner places PSC / spare-PSC by hand.
                    'psc': sum(b['psc'] for b in grp),
                    'sparePsc': sum(b['sparePsc'] for b in grp),
                    'backplanes': [b['code'] for b in grp],
                    'te': sum(b['te'] for b in grp),
                } for grp in groups],
                'actual': {
                    'racks': len(racks),
                    'backplanes': actual_bp,
                    'psc': sum(b['psc'] for r in racks for b in r['backplanes']),
                    'sparePsc': sum(b['sparePsc'] for r in racks for b in r['backplanes']),
                    'com': sum(b['com'] for r in racks for b in r['backplanes']),
                    'aeb': sum(b['aeb'] for r in racks for b in r['backplanes']),
                    'io': sum(b['io'] for r in racks for b in r['backplanes']),
                    'te': sum(b['te'] for r in racks for b in r['backplanes']),
                },
                'gesamt': gz,
                'rackTe': [sum(b['te'] for b in r['backplanes']) for r in racks],
            })
        out[proj] = locs
        wv.close()

    dest = os.path.join(PACKER, 'fixtures')
    os.makedirs(dest, exist_ok=True)
    with io.open(os.path.join(dest, 'reference.json'), 'w', encoding='utf-8') as fh:
        json.dump(out, fh, indent=2, ensure_ascii=False)

    # reconciliation report
    print('%-6s %-16s %-7s %-30s %s' % ('proj', 'location', 'groups', 'grid vs Gesamt', 'note'))
    print('-' * 96)
    bad = 0
    for proj, locs in out.items():
        for l in locs:
            a, gz = l['actual'], l['gesamt']
            checks = [('racks', a['racks'], gz['racks']), ('aeb', a['aeb'], gz['aeb']),
                      ('io', a['io'], gz['io']), ('psc', a['psc'], gz['psc']),
                      ('com', a['com'], gz['com'])]
            for code in ('BP-PWR-0', 'BP-PWR-4', 'BP-PWR-8', 'BP-EXB-1', 'BP-EXB-2', 'BP-EXB-4'):
                checks.append((code, a['backplanes'].get(code, 0), gz[code]))
            diff = ['%s %s!=%s' % (k, x, y) for k, x, y in checks if x != y]
            if diff:
                bad += 1
            print('%-6s %-16s %-7d %-30s %s'
                  % (proj, l['name'][:16], len(l['groups']),
                     'OK' if not diff else ', '.join(diff)[:30],
                     ' '.join('%dx%s' % (len(gr['backplanes']), gr['aeb']) for gr in l['groups'][:5])))
    print()
    print('locations: %d   mismatching: %d'
          % (sum(len(v) for v in out.values()), bad))


if __name__ == '__main__':
    main()
