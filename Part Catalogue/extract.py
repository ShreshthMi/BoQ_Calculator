# -*- coding: utf-8 -*-
"""Extract the BD BOM sheet into a structured part catalogue.

Source: BOM CAL/{ABS,Yard} V.1_2025-BRC with BD BOM.xlsm, sheet 'BD BOM', rows 4-144.
Both workbooks carry an identical catalogue in columns A-G; they differ only in the
per-location quantity block (H:BF) and its SUM in BG.
"""
import openpyxl, re, json, csv, os, warnings, unicodedata
from collections import Counter
warnings.filterwarnings('ignore')

SRC = {'ABS':  'BOM CAL/ABS V.1_2025-BRC with BD BOM.xlsm',
       'Yard': 'BOM CAL/Yard V.1_2025-BRC with BD BOM.xlsm'}
HANDOVER = 'BOM CAL/Handover BID Process Sheet Version 11.xlsx'
OUT = 'Part Catalogue'
ROW0, ROW1 = 4, 144          # catalogue rows
QC0, QC1   = 8, 58           # quantity columns H..BF
TOTAL_COL  = 59              # BG


def clean(v):
    if v is None:
        return ''
    s = unicodedata.normalize('NFKC', str(v)).replace(' ', ' ')
    return re.sub(r'\s+', ' ', s).strip()


def code(v):
    """Part numbers arrive as ints, floats or strings; normalise to a bare string."""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = clean(v)
    return '' if s in ('`', '-') else s


def split_markers(desc):
    """Strip the wrapping !, @, $ characters and report them separately.

    The workbook carries no legend for them, so they are preserved verbatim
    rather than interpreted.
    """
    core, found = desc, set()
    while True:
        m = re.match(r'^\s*([!@$`])\s*(.*)$', core)
        if m:
            found.add(m.group(1)); core = m.group(2); continue
        m = re.search(r'^(.*?)\s*([!@$`])\s*$', core)
        if m:
            found.add(m.group(2)); core = m.group(1); continue
        break
    return clean(core), ''.join(sorted(found))


# ---- Gesamt row labels, used to name each driver --------------------------
wb = openpyxl.load_workbook(SRC['ABS'], data_only=True)
gesamt = {r: clean(wb['Gesamt'].cell(r, 2).value) for r in range(1, 111)}
wb.close()

# ---- catalogue + per-project quantities -----------------------------------
parts, projects = {}, {}
for proj, path in SRC.items():
    wbf = openpyxl.load_workbook(path, data_only=False)   # formulas -> provenance
    wbv = openpyxl.load_workbook(path, data_only=True)    # cached values -> quantities
    wsf, wsv = wbf['BD BOM'], wbv['BD BOM']
    rev = clean(wbv['Revision'].cell(21, 2).value)

    locs = []
    for c in range(QC0, QC1 + 1):
        nm = clean(wsv.cell(3, c).value)
        used = bool(nm) and not re.fullmatch(r'(Tabelle|Loc)\s*\d+', nm)
        locs.append({'col': openpyxl.utils.get_column_letter(c),
                     'index': c - QC0 + 1, 'name': nm, 'used': used})
    used = [l for l in locs if l['used']]
    dup_labels = {}
    for l in used:
        dup_labels.setdefault(l['name'], []).append(l['col'])
    projects[proj] = {'workbook': os.path.basename(path), 'project': rev,
                      'locations': used,
                      'duplicate_location_labels': {k: v for k, v in dup_labels.items()
                                                    if len(v) > 1}}

    for r in range(ROW0, ROW1 + 1):
        raw = clean(wsv.cell(r, 6).value)
        if not raw:
            continue
        desc, markers = split_markers(raw)
        key = 'BD%03d' % r
        p = parts.setdefault(key, {
            'key': key, 'sl_no': wsv.cell(r, 1).value, 'source_row': r,
            'codes': {'at_sales_cloud': code(wsv.cell(r, 2).value),
                      'in_sales_cloud': code(wsv.cell(r, 3).value),
                      'ramco_erp':      code(wsv.cell(r, 4).value),
                      'rdso':           code(wsv.cell(r, 5).value)},
            'description': desc, 'description_raw': raw, 'markers': markers,
            'uom': clean(wsv.cell(r, 7).value), 'driver': None, 'quantities': {}})

        # provenance: what feeds the first location column
        if p['driver'] is None:
            f = wsf.cell(r, QC0).value
            if isinstance(f, str) and f.startswith('='):
                m = re.fullmatch(r'=Gesamt!\$?([A-Z]+)\$?(\d+)', f)
                if m:
                    gr = int(m.group(2))
                    p['driver'] = {'type': 'gesamt', 'ref': 'Gesamt!row %d' % gr,
                                   'label': gesamt.get(gr, ''), 'formula': f}
                else:
                    refs = sorted({'BD%03d' % int(n)
                                   for n in re.findall(r'\b[A-Z]{1,2}(\d+)\b', f)})
                    p['driver'] = {'type': 'derived_from_parts', 'ref': ', '.join(refs),
                                   'label': '', 'formula': f}
            else:
                p['driver'] = {'type': 'none', 'ref': '', 'label': '', 'formula': ''}

        # keyed by column, not by name: the Yard workbook reuses one location
        # label on two different columns, so a name-keyed map would lose a column
        per = []
        for l in locs:
            if not l['used']:
                continue
            v = wsv.cell(r, QC0 + l['index'] - 1).value
            if isinstance(v, (int, float)) and v:
                per.append({'col': l['col'], 'index': l['index'],
                            'location': l['name'], 'qty': v})
        tot = wsv.cell(r, TOTAL_COL).value
        p['quantities'][proj] = {
            'total': tot if isinstance(tot, (int, float)) else 0,
            'by_location': per}
    wbf.close(); wbv.close()

# ---- BoQ group from the handover sheet, matched on part number ------------
wb = openpyxl.load_workbook(HANDOVER, data_only=True)
ws = wb['10.  BOQ']
groups, cur = {}, ''
for r in range(3, 44):
    g = clean(ws.cell(r, 1).value)
    if g:
        cur = g.rstrip('.')
    c = code(ws.cell(r, 3).value)
    if c:
        groups[c] = cur
wb.close()
for p in parts.values():
    hit = [groups[c] for c in p['codes'].values() if c in groups]
    p['boq_group'] = hit[0] if hit else ''

# ---- duplicate part numbers ----------------------------------------------
seen = {}
for p in parts.values():
    for sysname, c in p['codes'].items():
        if c:
            seen.setdefault((sysname, c), []).append(p['key'])
dupes = {'%s:%s' % (s, c): ks for (s, c), ks in seen.items() if len(ks) > 1}
flagged = {k for ks in dupes.values() for k in ks}
for p in parts.values():
    p['duplicate_code'] = p['key'] in flagged

os.makedirs(OUT, exist_ok=True)
ordered = sorted(parts.values(), key=lambda p: p['source_row'])

# ---- parts.csv ------------------------------------------------------------
with open(os.path.join(OUT, 'parts.csv'), 'w', newline='', encoding='utf-8-sig') as fh:
    w = csv.writer(fh)
    w.writerow(['key', 'sl_no', 'source_row', 'boq_group', 'at_sales_cloud',
                'in_sales_cloud', 'ramco_erp', 'rdso', 'description', 'uom',
                'markers', 'duplicate_code', 'driver_type', 'driver_ref',
                'driver_label', 'driver_formula', 'qty_abs', 'qty_yard', 'qty_total'])
    for p in ordered:
        a = p['quantities'].get('ABS', {}).get('total', 0)
        y = p['quantities'].get('Yard', {}).get('total', 0)
        w.writerow([p['key'], p['sl_no'], p['source_row'], p['boq_group'],
                    p['codes']['at_sales_cloud'], p['codes']['in_sales_cloud'],
                    p['codes']['ramco_erp'], p['codes']['rdso'],
                    p['description'], p['uom'], p['markers'],
                    'yes' if p['duplicate_code'] else '',
                    p['driver']['type'], p['driver']['ref'],
                    p['driver']['label'], p['driver']['formula'], a, y, a + y])

# ---- quantities.csv (long form) -------------------------------------------
with open(os.path.join(OUT, 'quantities.csv'), 'w', newline='', encoding='utf-8-sig') as fh:
    w = csv.writer(fh)
    w.writerow(['key', 'description', 'project', 'location_col', 'location_index',
                'location', 'qty'])
    for p in ordered:
        for proj in ('ABS', 'Yard'):
            for e in p['quantities'].get(proj, {}).get('by_location', []):
                w.writerow([p['key'], p['description'], proj, e['col'],
                            e['index'], e['location'], e['qty']])

# ---- parts.json -----------------------------------------------------------
with open(os.path.join(OUT, 'parts.json'), 'w', encoding='utf-8') as fh:
    json.dump({'source': {'sheet': 'BD BOM', 'rows': '%d-%d' % (ROW0, ROW1),
                          'catalogue_version': 1.5, 'workbooks': projects,
                          'note': 'Columns A-G are identical in both workbooks; '
                                  'they differ only in the quantity block H:BF.'},
               'duplicate_codes': dupes,
               'parts': ordered}, fh, indent=2, ensure_ascii=False)

# ---- console summary ------------------------------------------------------
dt = Counter(p['driver']['type'] for p in ordered)
print('parts            : %d' % len(ordered))
print('driver types     : %s' % dict(dt))
print('with a UOM       : %d' % sum(1 for p in ordered if p['uom']))
print('grouped from BoQ : %d' % sum(1 for p in ordered if p['boq_group']))
print('marker-flagged   : %d' % sum(1 for p in ordered if p['markers']))
print('no part number   : %d' % sum(1 for p in ordered if not any(p['codes'].values())))
print('duplicate codes  : %d' % len(dupes))
for k, v in sorted(dupes.items()):
    print('   %-28s %s' % (k, ', '.join(v)))
print('used on project  : %d' % sum(1 for p in ordered
                                    if p['quantities']['ABS']['total']
                                    or p['quantities']['Yard']['total']))
for k, v in projects.items():
    print('  %-5s %-24r %d locations  dup labels: %s'
          % (k, v['project'], len(v['locations']),
             v['duplicate_location_labels'] or 'none'))
